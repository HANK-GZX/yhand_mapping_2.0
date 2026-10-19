#!/usr/bin/env python3
"""Plot distributions of paired result/glove_data and result/motor_feedback.

Each point is an average, not an individual frame. Only the ``data`` worksheet
is read; point_log.csv supplies combo labels and collection status.
"""

import argparse
import csv
import json
import math
import re
from collections import Counter
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

from hand_motor_map import HAND_MOTOR_MAP, MOTOR_COUNT


FINGERS = ("thumb", "index", "middle", "ring", "little")
FIELDS = ("pos", "vel", "tor", "kp", "kd")
COLORS = {True: "#2878b5", False: "#d9534f"}


def arguments():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("result_dir", type=Path, help="Single finger's result directory or a hand collection directory")
    parser.add_argument("--output-dir", type=Path, help="Default: input_dir/combo_distribution")
    parser.add_argument("--point-log", type=Path, help="Default: result_dir/../point_log.csv")
    parser.add_argument("--hand", choices=("left", "right"), help="Normally inferred from path")
    parser.add_argument("--finger", choices=FINGERS, help="Normally inferred from path")
    parser.add_argument("--fingers", nargs="+", choices=FINGERS,
                        help="Select fingers when input is a hand collection directory")
    parser.add_argument("--combo", nargs="+", help="Plot only these combo IDs")
    parser.add_argument("--reached-only", action="store_true", help="Keep reached=1, timeout=0")
    parser.add_argument("--motor-fields", nargs="+", choices=FIELDS, default=list(FIELDS[:3]))
    parser.add_argument("--bins", type=int, default=40)
    parser.add_argument("--dpi", type=int, default=150)
    args = parser.parse_args()
    if args.bins < 2 or args.dpi < 1:
        parser.error("--bins must be >=2 and --dpi must be positive")
    if args.finger and args.fingers:
        parser.error("Use either --finger or --fingers")
    return args


def numeric_array(rows, shape, label):
    if len(rows) != shape[0] or any(len(row) != shape[1] for row in rows):
        raise ValueError(f"{label}: expected shape {shape}")
    # Empty Excel cells represent missing values, not zeros.
    result = np.asarray([[np.nan if v is None else float(v) for v in row] for row in rows])
    if np.isinf(result).any():
        raise ValueError(f"{label}: contains infinity")
    return result


def read_data(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "data" not in workbook.sheetnames:
            raise ValueError(f"{path.name}: missing data worksheet")
        return list(workbook["data"].values)
    finally:
        workbook.close()


def file_index(folder):
    result = {}
    for path in sorted(folder.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        if not path.stem.isdigit():
            raise ValueError(f"Unexpected point filename: {path}")
        point_id = int(path.stem)
        if point_id in result:
            raise ValueError(f"Duplicate point ID {point_id} in {folder}")
        result[point_id] = path
    if not result:
        raise ValueError(f"No point XLSX files in {folder}")
    return result


def vector_repetition(points, key):
    valid = [(p["point_id"], tuple(p[key])) for p in points if np.isfinite(p[key]).all()]
    counts = Counter(vector for _, vector in valid)
    if not counts:
        return {"finite_vector_count": 0, "unique_vector_count": 0}
    vector, count = counts.most_common(1)[0]
    longest = {"point_count": 0, "first_point_id": None, "last_point_id": None}
    previous_vector, previous_id, run_start, run_length = None, None, None, 0
    for point_id, current in valid:
        if current == previous_vector and point_id == previous_id + 1:
            run_length += 1
        else:
            run_start, run_length = point_id, 1
        if run_length > longest["point_count"]:
            longest = {"point_count": run_length, "first_point_id": run_start, "last_point_id": point_id}
        previous_vector, previous_id = current, point_id
    return {"finite_vector_count": len(valid), "unique_vector_count": len(counts),
            "most_repeated_vector": list(vector), "most_repeated_count": count,
            "longest_identical_run": longest}


def load_points(result_dir, point_log, hand, finger, combo_filter, reached_only):
    glove_files = file_index(result_dir / "glove_data")
    motor_files = file_index(result_dir / "motor_feedback")
    metadata = {}
    with point_log.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        required = {"point_id", "finger", "combo_id", "reached", "timeout", "active_mode"}
        if not required.issubset(reader.fieldnames or []):
            raise ValueError(f"point_log.csv missing columns: {sorted(required - set(reader.fieldnames or []))}")
        for row in reader:
            point_id = int(row["point_id"])
            if row["finger"] != finger:
                raise ValueError(f"Point {point_id}: finger mismatch in point log")
            if point_id in metadata:
                raise ValueError(f"Duplicate point ID {point_id} in point log")
            if row["reached"] not in ("0", "1") or row["timeout"] not in ("0", "1"):
                raise ValueError(f"Point {point_id}: invalid reached/timeout flag")
            metadata[point_id] = row
    requested = set(combo_filter or [])
    unknown = requested - {row["combo_id"] for row in metadata.values()}
    if unknown:
        raise ValueError(f"Unknown combos: {sorted(unknown)}")
    report = {
        "result_dir": str(result_dir), "point_log": str(point_log),
        "hand": hand, "finger": finger, "reached_only": reached_only,
        "combo_filter": sorted(requested),
        "glove_file_count": len(glove_files), "motor_file_count": len(motor_files),
        "point_log_count": len(metadata),
        "missing_glove_point_ids": sorted((set(metadata) | set(motor_files)) - set(glove_files)),
        "missing_motor_point_ids": sorted((set(metadata) | set(glove_files)) - set(motor_files)),
        "unlabelled_point_ids": sorted((set(glove_files) | set(motor_files)) - set(metadata)),
        "invalid_points": [],
    }
    selected = [pid for pid in sorted(set(metadata) & set(glove_files) & set(motor_files))
                if (not requested or metadata[pid]["combo_id"] in requested)
                and (not reached_only or (metadata[pid]["reached"] == "1" and metadata[pid]["timeout"] == "0"))]
    points = []
    q_offset = (8 if hand == "left" else 28) + 4 * FINGERS.index(finger)
    for number, point_id in enumerate(selected, 1):
        try:
            glove = numeric_array(read_data(glove_files[point_id]), (1, 48), "glove")
            rows = read_data(motor_files[point_id])
            if len(rows) != MOTOR_COUNT + 1 or len(rows[0]) != 5:
                raise ValueError("motor_feedback: expected header plus 28 x 5 matrix")
            window, sample_count = rows[0][:2]
            if not isinstance(window, str) or not window:
                raise ValueError("motor_feedback: invalid averaging-window header")
            if sample_count is None or float(sample_count) <= 0 or not float(sample_count).is_integer():
                raise ValueError("motor_feedback: invalid sample_count")
            motor = numeric_array(rows[1:], (MOTOR_COUNT, 5), "motor_feedback")
            points.append({"point_id": point_id, "meta": metadata[point_id],
                           "glove": glove[0], "q": glove[0, q_offset:q_offset + 4],
                           "motor": motor, "window": window, "sample_count": int(sample_count)})
        except Exception as exc:
            report["invalid_points"].append({"point_id": point_id, "error": str(exc)})
        if number % 500 == 0 or number == len(selected):
            print(f"Loaded {number}/{len(selected)} pairs; invalid={len(report['invalid_points'])}", flush=True)
    report["selected_pair_count"] = len(selected)
    report["loaded_point_count"] = len(points)
    report["glove_q_columns_1_based"] = list(range(q_offset + 1, q_offset + 5))
    report["averaging_windows"] = dict(Counter(p["window"] for p in points))
    report["glove_q_repetition"] = vector_repetition(points, "q")
    report["glove48_repetition"] = vector_repetition(points, "glove")
    return points, report


def combo_axes(combo, hand, finger):
    """Interpret the collector's combo names using the repository's wiring map."""
    pattern = rf"{finger}_(.+)_(in|out)(?:_opp_(?:on|off))?"
    match = re.fullmatch(pattern, combo)
    if match is None:
        raise ValueError(f"Unsupported combo name: {combo}")
    role, direction = match.groups()
    role = {"deep": "flex_deep", "shallow": "flex_shallow", "flex": "flexor", "ext": "extensor"}.get(role, role)
    motors = HAND_MOTOR_MAP[hand][finger]
    return motors[role], motors["adduction" if direction == "in" else "abduction"]


def save_plot(fig, path, dpi):
    fig.savefig(path, dpi=dpi)
    plt.close(fig)


def finite_limits(values):
    values = values[np.isfinite(values)]
    if not values.size:
        return (-1, 1)
    low, high = float(values.min()), float(values.max())
    pad = max((high - low) * 0.04, 0.01)
    return low - pad, high + pad


def quantile_limits(values):
    finite = values[np.isfinite(values)]
    if not finite.size:
        return (-1, 1)
    return finite_limits(np.percentile(finite, [1, 99]))


def histogram(ax, values, reached, edges):
    for status, label in ((True, "reached, no timeout"), (False, "unreached / timeout")):
        subset = values[(reached == status) & np.isfinite(values)]
        if subset.size:
            ax.hist(subset, bins=edges, histtype="step", linewidth=1.5,
                    color=COLORS[status], label=label)
    ax.set_ylabel("Point count")
    ax.grid(alpha=0.2)


def write_csv(path, rows):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def plot_distributions(points, output, hand, finger, fields, bins, dpi):
    motor_names = {mid: role for role, mid in HAND_MOTOR_MAP[hand][finger].items()}
    motor_ids = sorted(motor_names)
    combos = sorted({p["meta"]["combo_id"] for p in points})
    all_motor = np.stack([p["motor"] for p in points])
    all_q = np.stack([p["q"] for p in points])
    # Position panels show combo axes. Exclude a motor's non-axis positions
    # from these shared limits; its full distribution stays in the histograms.
    pos_limits = {mid: finite_limits(np.array([p["motor"][mid - 1, 0] for p in points
                  if mid in combo_axes(p["meta"]["combo_id"], hand, finger)])) for mid in motor_ids}
    q_limits = [finite_limits(all_q[:, qi]) for qi in range(4)]
    # Shared bins/limits across combos make the plots directly comparable.
    motor_edges = {(mid, field): np.linspace(*finite_limits(all_motor[:, mid - 1, FIELDS.index(field)]), bins + 1)
                   for mid in motor_ids for field in fields}
    q_edges = [np.linspace(*limits, bins + 1) for limits in q_limits]
    overview, axes = plt.subplots(math.ceil(len(combos) / 3), 3, figsize=(16, 4.5 * math.ceil(len(combos) / 3)),
                                  squeeze=False, layout="constrained")
    zoom, zoom_axes = plt.subplots(math.ceil(len(combos) / 3), 3,
                                  figsize=(16, 4.5 * math.ceil(len(combos) / 3)),
                                  squeeze=False, layout="constrained")
    summaries, statistics, point_rows = [], [], []
    for combo, overview_ax in zip(combos, axes.flat):
        group = [p for p in points if p["meta"]["combo_id"] == combo]
        motor = np.stack([p["motor"] for p in group])
        q = np.stack([p["q"] for p in group])
        ids = np.asarray([p["point_id"] for p in group])
        reached = np.asarray([p["meta"]["reached"] == "1" and p["meta"]["timeout"] == "0" for p in group])
        x_motor, y_motor = combo_axes(combo, hand, finger)
        x, y = motor[:, x_motor - 1, 0], motor[:, y_motor - 1, 0]
        x_label, y_label = f"M{x_motor} {motor_names[x_motor]} pos", f"M{y_motor} {motor_names[y_motor]} pos"
        for status, marker in ((True, "o"), (False, "x")):
            mask = reached == status
            overview_ax.scatter(x[mask], y[mask], s=9, alpha=0.45, marker=marker,
                                color=COLORS[status], linewidths=0.5,
                                label="reached, no timeout" if status else "unreached / timeout")
        overview_ax.set(title=f"{combo} | n={len(group)} | timeout={sum(p['meta']['timeout'] == '1' for p in group)}",
                        xlabel=x_label, ylabel=y_label, xlim=pos_limits[x_motor], ylim=pos_limits[y_motor])
        overview_ax.legend(fontsize=7)
        overview_ax.grid(alpha=0.2)
        zoom_ax = zoom_axes.flat[combos.index(combo)]
        x_zoom, y_zoom = quantile_limits(x), quantile_limits(y)
        outside = np.isfinite(x) & np.isfinite(y) & (
            (x < x_zoom[0]) | (x > x_zoom[1]) | (y < y_zoom[0]) | (y > y_zoom[1]))
        for status, marker in ((True, "o"), (False, "x")):
            mask = reached == status
            zoom_ax.scatter(x[mask], y[mask], s=9, alpha=0.45, marker=marker,
                            color=COLORS[status], linewidths=0.5,
                            label="reached, no timeout" if status else "unreached / timeout")
        zoom_ax.set(title=f"{combo} | n={len(group)} | outside zoom={int(outside.sum())}",
                    xlabel=x_label, ylabel=y_label, xlim=x_zoom, ylim=y_zoom)
        zoom_ax.legend(fontsize=7)
        zoom_ax.grid(alpha=0.2)

        folder = output / combo
        folder.mkdir(parents=True, exist_ok=True)
        fig, axs = plt.subplots(2, 3, figsize=(16, 9), layout="constrained")
        valid_xy = np.isfinite(x) & np.isfinite(y)
        artist = axs.flat[0].scatter(x[valid_xy], y[valid_xy], c=ids[valid_xy], s=8, alpha=0.7, cmap="viridis")
        fig.colorbar(artist, ax=axs.flat[0], label="Point ID (collection order)")
        axs.flat[0].set_title("Motor positions by collection order")
        if valid_xy.any():
            artist = axs.flat[1].hexbin(x[valid_xy], y[valid_xy], gridsize=bins, mincnt=1, cmap="magma",
                                       extent=(*pos_limits[x_motor], *pos_limits[y_motor]))
            fig.colorbar(artist, ax=axs.flat[1], label="Point count")
        axs.flat[1].set_title("Motor position density")
        for qi, ax in enumerate(list(axs.flat)[2:]):
            valid = valid_xy & np.isfinite(q[:, qi])
            artist = ax.scatter(x[valid], y[valid], c=q[valid, qi], s=9, alpha=0.7, cmap="viridis",
                                vmin=q_limits[qi][0], vmax=q_limits[qi][1])
            fig.colorbar(artist, ax=ax, label=f"glove q{qi}")
            ax.set_title(f"Motor positions colored by glove q{qi}")
        for ax in axs.flat:
            ax.set(xlabel=x_label, ylabel=y_label, xlim=pos_limits[x_motor], ylim=pos_limits[y_motor])
            ax.grid(alpha=0.2)
        fig.suptitle(f"{combo} | {len(group)} averaged points")
        save_plot(fig, folder / "motor_position_distribution.png", dpi)

        fig, axs = plt.subplots(len(motor_ids), len(fields), figsize=(5 * len(fields), 2.7 * len(motor_ids)),
                                squeeze=False, layout="constrained")
        for row, mid in enumerate(motor_ids):
            for col, field in enumerate(fields):
                ax = axs[row, col]
                values = motor[:, mid - 1, FIELDS.index(field)]
                histogram(ax, values, reached, motor_edges[mid, field])
                ax.set(title=f"M{mid} {motor_names[mid]} | {field}", xlabel=f"{field} (recorded units)")
        axs[0, 0].legend(fontsize=8)
        fig.suptitle(f"{combo} | feedback distributions for all {finger} motors")
        save_plot(fig, folder / "motor_feedback_histograms.png", dpi)

        fig, axs = plt.subplots(2, 4, figsize=(17, 7), layout="constrained")
        for qi in range(4):
            histogram(axs[0, qi], q[:, qi], reached, q_edges[qi])
            axs[0, qi].set(title=f"glove q{qi} distribution", xlabel=f"glove q{qi}")
            for status in (True, False):
                mask = reached == status
                axs[1, qi].scatter(ids[mask], q[mask, qi], color=COLORS[status], s=7, alpha=0.5)
            axs[1, qi].set(title=f"glove q{qi} vs collection order", xlabel="Point ID", ylabel=f"glove q{qi}",
                           ylim=q_limits[qi])
            axs[1, qi].grid(alpha=0.2)
        axs[0, 0].legend(fontsize=8)
        fig.suptitle(f"{combo} | glove distributions and changes during collection")
        save_plot(fig, folder / "glove_distribution.png", dpi)

        summaries.append({"combo_id": combo, "point_count": len(group), "reached_no_timeout_count": int(reached.sum()),
                          "logged_reached_count": sum(p["meta"]["reached"] == "1" for p in group),
                          "timeout_count": sum(p["meta"]["timeout"] == "1" for p in group),
                          "x_motor": x_motor, "y_motor": y_motor,
                          "x_pos_missing_count": int((~np.isfinite(x)).sum()),
                          "y_pos_missing_count": int((~np.isfinite(y)).sum()),
                          "active_mode_counts": json.dumps(dict(Counter(p["meta"]["active_mode"] for p in group)), sort_keys=True)})
        channels = [(f"motor_{mid}_{field}", motor[:, mid - 1, fi])
                    for mid in motor_ids for fi, field in enumerate(FIELDS)]
        channels += [(f"glove_q{qi}", q[:, qi]) for qi in range(4)]
        for name, values in channels:
            finite = values[np.isfinite(values)]
            stats = {"combo_id": combo, "field": name, "finite_count": int(finite.size),
                     "missing_count": int(len(values) - finite.size)}
            stats.update(dict(zip(("min", "p05", "median", "p95", "max"),
                                  np.percentile(finite, [0, 5, 50, 95, 100]) if finite.size else [None] * 5)))
            stats.update(mean=float(finite.mean()) if finite.size else None,
                         std=float(finite.std()) if finite.size else None)
            statistics.append(stats)
        for p in group:
            record = {"point_id": p["point_id"], "combo_id": combo, "reached": p["meta"]["reached"],
                      "timeout": p["meta"]["timeout"], "active_mode": p["meta"]["active_mode"],
                      "window": p["window"], "sample_count": p["sample_count"]}
            record.update({f"glove_q{qi}": p["q"][qi] for qi in range(4)})
            record.update({f"motor_{mid}_{field}": p["motor"][mid - 1, fi]
                           for mid in motor_ids for fi, field in enumerate(FIELDS)})
            point_rows.append(record)
        print(f"Plotted {combo}: n={len(group)}, reached/no-timeout={int(reached.sum())}", flush=True)
    for ax in list(axes.flat)[len(combos):]:
        ax.set_visible(False)
    for ax in list(zoom_axes.flat)[len(combos):]:
        ax.set_visible(False)
    overview.suptitle(f"{hand} {finger} | measured motor positions | all plotted points are result-window averages")
    save_plot(overview, output / "combo_overview.png", dpi)
    zoom.suptitle(f"{hand} {finger} | per-combo P01-P99 zoom | axes vary by panel; all points remain in exports")
    save_plot(zoom, output / "combo_overview_zoom.png", dpi)
    write_csv(output / "combo_summary.csv", summaries)
    write_csv(output / "field_statistics.csv", statistics)
    write_csv(output / "points.csv", sorted(point_rows, key=lambda row: row["point_id"]))
    return summaries


def run_result(args, result_dir, output, hand, finger, combo_filter):
    point_log = (args.point_log or result_dir.parent / "point_log.csv").expanduser().resolve()
    points, report = load_points(result_dir, point_log, hand, finger, combo_filter, args.reached_only)
    output.mkdir(parents=True, exist_ok=True)
    report_path = output / "load_report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if not points:
        raise ValueError(f"No valid selected points; see {report_path}")
    report["combos"] = plot_distributions(points, output, hand, finger, args.motor_fields, args.bins, args.dpi)
    report["notes"] = ["Statistics describe distributions across averaged collection points, not temporal variability within a point.",
                       "Default includes unreached/timeout points; flags come directly from point_log.csv.",
                       "Motor IDs are interpreted using hand_motor_map.py; values remain in recorded units.",
                       "Histograms, position axes, and glove color scales share limits across selected combos.",
                       "Shared position limits use only combos where each motor is an x/y axis; histograms include its full distribution.",
                       "combo_overview_zoom.png uses per-combo P01-P99 limits plus padding and labels the number of points outside each view.",
                       "glove repetition counts use exact finite values and consecutive numeric point IDs; repeated data is retained.",
                       "std is population standard deviation (ddof=0); non-finite cells are counted as missing."]
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Output: {output}")
    for key in ("missing_glove_point_ids", "missing_motor_point_ids", "unlabelled_point_ids", "invalid_points"):
        if report[key]:
            print(f"WARNING: {key}={len(report[key])}; see load_report.json")
    return report


def main():
    args = arguments()
    input_dir = args.result_dir.expanduser().resolve()
    hand = args.hand or next((side for side in ("left", "right") if f"{side}_hand" in input_dir.parts), None)
    if hand is None:
        raise ValueError("Cannot infer hand; provide --hand")
    output = (args.output_dir or input_dir / "combo_distribution").expanduser().resolve()
    if (input_dir / "glove_data").is_dir() and (input_dir / "motor_feedback").is_dir():
        if args.fingers:
            raise ValueError("--fingers requires a hand collection directory")
        finger = args.finger or input_dir.parent.name
        if finger not in FINGERS:
            raise ValueError("Cannot infer finger; provide --finger")
        run_result(args, input_dir, output, hand, finger, args.combo)
        return

    if args.point_log:
        raise ValueError("--point-log is only supported for a single result directory")
    available = [finger for finger in FINGERS
                 if (input_dir / finger / "result" / "glove_data").is_dir()
                 and (input_dir / finger / "result" / "motor_feedback").is_dir()]
    selected = set(args.fingers or ([args.finger] if args.finger else available))
    missing = selected - set(available)
    if missing:
        raise ValueError(f"Missing finger result directories: {sorted(missing)}")
    if not available:
        raise ValueError(f"No finger result directories under {input_dir}")
    jobs = [finger for finger in available if finger in selected]
    if args.combo:
        unknown = [combo for combo in args.combo
                   if not any(combo.startswith(f"{finger}_") for finger in jobs)]
        if unknown:
            raise ValueError(f"Combos do not belong to selected fingers: {unknown}")
        jobs = [finger for finger in jobs if any(combo.startswith(f"{finger}_") for combo in args.combo)]
    reports, combo_rows, finger_rows = [], [], []
    for finger in jobs:
        combo_filter = [combo for combo in args.combo if combo.startswith(f"{finger}_")] if args.combo else None
        print(f"Analyzing {hand} {finger}", flush=True)
        report = run_result(args, input_dir / finger / "result", output / finger, hand, finger, combo_filter)
        reports.append(report)
        combo_rows.extend({"hand": hand, "finger": finger, **row} for row in report["combos"])
        total = report["loaded_point_count"]
        reached = sum(row["reached_no_timeout_count"] for row in report["combos"])
        timeout = sum(row["timeout_count"] for row in report["combos"])
        finger_rows.append({"hand": hand, "finger": finger, "point_count": total,
                            "combo_count": len(report["combos"]),
                            "reached_no_timeout_count": reached, "timeout_count": timeout,
                            "timeout_ratio": timeout / total,
                            "output_dir": str(output / finger)})
    write_csv(output / "collection_combo_summary.csv", combo_rows)
    write_csv(output / "collection_finger_summary.csv", finger_rows)
    collection_report = {"input_dir": str(input_dir), "hand": hand,
                         "point_count": sum(row["point_count"] for row in finger_rows),
                         "fingers": finger_rows,
                         "load_reports": [str(output / r["finger"] / "load_report.json") for r in reports]}
    (output / "collection_report.json").write_text(json.dumps(collection_report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Collection output: {output}", flush=True)


if __name__ == "__main__":
    main()
